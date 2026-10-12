"""
AstroMobbin Quiz — Flask Backend
Serves the web app and persists quiz results to results/quiz_results.xlsx
"""

import os
import json
import traceback
from datetime import datetime
from pathlib import Path

# ── Attempt to import flask & openpyxl ───────────────────────
try:
    from flask import Flask, request, jsonify, send_from_directory, abort
except ImportError:
    import subprocess, sys
    print("Installing Flask...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "flask", "--quiet"])
    from flask import Flask, request, jsonify, send_from_directory, abort

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess, sys
    print("Installing openpyxl...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "--quiet"])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

# ── Flask App ─────────────────────────────────────────────────
BASE_DIR    = Path(__file__).parent
RESULTS_DIR = BASE_DIR / "results"
XLSX_PATH   = RESULTS_DIR / "quiz_results.xlsx"

app = Flask(__name__, static_folder=str(BASE_DIR), static_url_path="")

# ── Excel helpers ──────────────────────────────────────────────
HEADERS = [
    "Trial",
    "Timestamp",
    "Name",
    "Department",
    "ID (Last 5 Digits)",
    "Score",
    "Total Questions",
    "Percentage (%)",
    "Timer (sec/q)",
    "Time Taken (sec)"
]

def get_or_create_workbook():
    """Open existing workbook or create a styled blank one."""
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    if XLSX_PATH.exists():
        return openpyxl.load_workbook(XLSX_PATH)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quiz Results"

    # Styling constants
    header_fill  = PatternFill("solid", fgColor="0E1422")
    header_font  = Font(bold=True, color="38BDF8", name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border  = Border(
        left   = Side(style="thin", color="334155"),
        right  = Side(style="thin", color="334155"),
        top    = Side(style="thin", color="334155"),
        bottom = Side(style="thin", color="334155")
    )

    # Write + style header row
    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill      = header_fill
        cell.font      = header_font
        cell.alignment = center_align
        cell.border    = thin_border

    # Column widths
    col_widths = [8, 22, 28, 22, 20, 8, 16, 16, 16, 18]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(XLSX_PATH)
    return wb


def get_trial_number(ws, name: str, id5: str) -> int:
    """Count how many rows exist for this name+id5 pair and return next trial number."""
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[2] and row[4]:  # Name col=3 (idx 2), ID col=5 (idx 4)
            if (str(row[2]).strip().lower() == name.strip().lower() and
                    str(row[4]).strip() == id5.strip()):
                count += 1
    return count + 1   # next trial number


def append_result(payload: dict) -> int:
    """Append a result row to the xlsx and return the trial number."""
    wb = get_or_create_workbook()
    ws = wb.active

    trial = get_trial_number(ws, payload["name"], payload["id5"])
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Alternating row colour
    next_row = ws.max_row + 1
    row_fill = PatternFill("solid", fgColor="141E30" if next_row % 2 == 0 else "0E1422")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border  = Border(
        left   = Side(style="thin", color="334155"),
        right  = Side(style="thin", color="334155"),
        top    = Side(style="thin", color="334155"),
        bottom = Side(style="thin", color="334155")
    )

    # Colour-code score percentage in the percentage cell
    pct = payload.get("percentage", 0)
    pct_color = "22D3A0" if pct >= 80 else ("FBBF24" if pct >= 60 else "F87171")

    row_data = [
        trial,
        timestamp,
        payload["name"],
        payload["department"],
        payload["id5"],
        payload["score"],
        payload["total"],
        payload["percentage"],
        payload["timer_secs"],
        payload["time_taken"]
    ]

    for col_idx, value in enumerate(row_data, start=1):
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.fill      = row_fill
        cell.alignment = center_align
        cell.border    = thin_border

        # Percentage column (8) — colour-coded
        if col_idx == 8:
            cell.font = Font(bold=True, color=pct_color, name="Calibri")
        # Trial column (1) — bold
        elif col_idx == 1:
            cell.font = Font(bold=True, name="Calibri")

    ws.row_dimensions[next_row].height = 22
    wb.save(XLSX_PATH)
    return trial


# ── Routes ─────────────────────────────────────────────────────

@app.route("/")
def index():
    """Serve the main HTML page."""
    return send_from_directory(str(BASE_DIR), "index.html")


@app.route("/<path:filename>")
def static_files(filename):
    """Serve CSS, JS, and other static files."""
    try:
        return send_from_directory(str(BASE_DIR), filename)
    except Exception:
        abort(404)


@app.route("/api/save-result", methods=["POST"])
def save_result():
    """
    POST JSON body:
      name, department, id5, score, total, percentage, timer_secs, time_taken

    Returns:
      { "status": "ok", "trial": <int>, "message": "..." }
    """
    try:
        data = request.get_json(force=True)

        required = ["name", "department", "id5", "score", "total",
                    "percentage", "timer_secs", "time_taken"]
        for field in required:
            if field not in data:
                return jsonify(status="error", message=f"Missing field: {field}"), 400

        # Basic type coercion
        payload = {
            "name":       str(data["name"]).strip(),
            "department": str(data["department"]).strip(),
            "id5":        str(data["id5"]).strip(),
            "score":      int(data["score"]),
            "total":      int(data["total"]),
            "percentage": int(data["percentage"]),
            "timer_secs": int(data["timer_secs"]),
            "time_taken": int(data["time_taken"])
        }

        trial = append_result(payload)

        print(f"[SAVED] {payload['name']} | {payload['department']} | "
              f"Score: {payload['score']}/{payload['total']} | Trial: {trial}")

        return jsonify(
            status="ok",
            trial=trial,
            message=f"Result saved (Trial #{trial})"
        )

    except Exception as e:
        traceback.print_exc()
        return jsonify(status="error", message=str(e)), 500


# ── Entry Point ────────────────────────────────────────────────
if __name__ == "__main__":
    import webbrowser
    PORT = 5000

    # Find available port
    import socket
    for port in range(PORT, PORT + 20):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(("localhost", port)) != 0:
                PORT = port
                break

    print("=" * 58)
    print("  ASTROMOBBIN Space Quiz Platform")
    print(f"  Running at: http://localhost:{PORT}")
    print(f"  Results Excel: {XLSX_PATH}")
    print("  Press Ctrl+C to stop")
    print("=" * 58)

    try:
        webbrowser.open(f"http://localhost:{PORT}")
    except Exception:
        pass

    app.run(host="0.0.0.0", port=PORT, debug=False)
